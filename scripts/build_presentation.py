"""Build the AgroVision hackathon pitch deck from project artifacts."""

# ruff: noqa: E501

from __future__ import annotations

from pathlib import Path
from typing import Final

from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT: Final = Path(__file__).resolve().parents[1]
OUT_DIR: Final = ROOT / "docs" / "presentation"
ASSETS: Final = OUT_DIR / "assets"
SCREENS: Final = ASSETS / "screenshots"
GENERATED: Final = ASSETS / "generated"
DECK_PATH: Final = OUT_DIR / "AgroVision_pitch_deck.pptx"

NAVY = "07111F"
NAVY_2 = "0C1B2A"
INK = "102234"
MUTED = "65758B"
WHITE = "FFFFFF"
MIST = "F4FAF7"
CARD = "FFFFFF"
GREEN = "18A66A"
GREEN_2 = "4ADE80"
LIME = "C9F45A"
CYAN = "62D5D2"
AMBER = "F6B73C"
ROSE = "F06472"
GRID = "DDEBE5"

W = Inches(13.333)
H = Inches(7.5)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_rect(slide, x, y, w, h, fill, radius=True, line=None, transparency=0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=20,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, parts, x, y, w, h, size=20, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, part_color, bold in parts:
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos Display"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(part_color or color)
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=INK, accent=GREEN):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, item in enumerate(items):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        paragraph.line_spacing = 1.06
        lead = paragraph.add_run()
        lead.text = "●  "
        lead.font.name = "Aptos"
        lead.font.size = Pt(size - 2)
        lead.font.color.rgb = rgb(accent)
        run = paragraph.add_run()
        run.text = item
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
    return box


def add_picture_cover(slide, path: Path, x, y, w, h, brightness=1.0):
    path = Path(path)
    source = path
    if brightness != 1.0:
        GENERATED.mkdir(parents=True, exist_ok=True)
        source = GENERATED / f"{path.stem}-b{brightness:.2f}.jpg"
        if not source.exists():
            image = Image.open(path).convert("RGB")
            ImageEnhance.Brightness(image).enhance(brightness).save(source, quality=92)
    with Image.open(source) as image:
        iw, ih = image.size
    frame_ratio = w / h
    image_ratio = iw / ih
    picture = slide.shapes.add_picture(str(source), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        crop = (1 - image_ratio / frame_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_picture_contain(slide, path: Path, x, y, w, h):
    with Image.open(path) as image:
        iw, ih = image.size
    ratio = min(w / iw, h / ih)
    width, height = iw * ratio, ih * ratio
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - width) / 2), Inches(y + (h - height) / 2), Inches(width), Inches(height)
    )


def add_frame(slide, x, y, w, h, color=WHITE, line="D6E8E0"):
    add_rect(slide, x + 0.05, y + 0.10, w, h, NAVY, transparency=82, line=NAVY)
    return add_rect(slide, x, y, w, h, color, line=line)


def add_pill(slide, text, x, y, w, fill=LIME, color=NAVY, size=10):
    add_rect(slide, x, y, w, 0.34, fill)
    add_text(slide, text.upper(), x, y + 0.01, w, 0.29, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_title(slide, kicker, title, subtitle=None, dark=False):
    base = WHITE if dark else INK
    add_text(slide, kicker.upper(), 0.65, 0.35, 5.7, 0.25, size=9, color=GREEN_2 if dark else GREEN, bold=True)
    add_text(slide, title, 0.65, 0.69, 12.0, 0.72, size=28, color=base, bold=True, font="Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.67, 1.40, 11.8, 0.42, size=12, color="B6C7D7" if dark else MUTED)


def add_footer(slide, number, dark=False, source=None):
    color = "8DA1B3" if dark else "8A9B98"
    if source:
        add_text(slide, source, 0.66, 7.16, 10.7, 0.18, size=7.5, color=color)
    add_text(slide, f"AGROVISION   /   {number:02d}", 11.33, 7.12, 1.35, 0.20, size=8, color=color, bold=True, align=PP_ALIGN.RIGHT)


def blank(prs, color=MIST):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(color)
    return slide


def add_metric(slide, value, label, x, y, w, color=GREEN, suffix=""):
    add_rect(slide, x, y, w, 1.07, CARD, line="D9EBE3")
    add_rich_text(
        slide,
        [(value, color, True), (suffix, color, True)],
        x + 0.22,
        y + 0.16,
        w - 0.44,
        0.43,
        size=24,
    )
    add_text(slide, label, x + 0.22, y + 0.69, w - 0.44, 0.22, size=9, color=MUTED, bold=True)


def connect(slide, x1, y1, x2, y2, color=GREEN, width=2.0):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_node(slide, title, subtitle, x, y, w, h, fill=WHITE, accent=GREEN, dark=False):
    add_rect(slide, x, y, w, h, fill, line=accent)
    add_rect(slide, x + 0.16, y + 0.17, 0.10, h - 0.34, accent, radius=False)
    add_text(slide, title, x + 0.38, y + 0.18, w - 0.55, 0.28, size=13, color=WHITE if dark else INK, bold=True)
    add_text(slide, subtitle, x + 0.38, y + 0.55, w - 0.55, h - 0.66, size=9.5, color="B6C7D7" if dark else MUTED)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "AgroVision — интеллектуальный мониторинг поголовья"
    prs.core_properties.subject = "Hackathon pitch deck"
    prs.core_properties.author = "Команда AgroVision"
    prs.core_properties.comments = "Generated from verified project artifacts."

    wide = ROOT / "artifacts/demo_predictions/posters/7438342-uhd_4096_1974_30fps.jpg"
    tall = ROOT / "artifacts/demo_predictions/posters/12831053_1080_1920_60fps.jpg"
    dashboard = SCREENS / "02-dashboard-top.png"
    controls = SCREENS / "03-dashboard-controls.png"
    upload = SCREENS / "05-upload.png"
    test_hero = ROOT / "artifacts/demo_predictions/images/example-03.jpg"
    best_video = ROOT / "artifacts/demo_predictions/videos/7438342-uhd_4096_1974_30fps.mp4"
    examples = [ROOT / f"artifacts/demo_predictions/images/example-{i:02d}.jpg" for i in (2, 4, 6)]

    # 01 — Cover
    slide = blank(prs, NAVY)
    add_picture_cover(slide, wide, 0, 0, 13.333, 7.5, brightness=0.47)
    add_rect(slide, 0, 0, 13.333, 7.5, NAVY, radius=False, transparency=28)
    add_rect(slide, 0, 0, 7.55, 7.5, NAVY, radius=False, transparency=7)
    add_rect(slide, 7.70, 0.66, 4.96, 6.18, NAVY_2, line="294457")
    add_picture_cover(slide, wide, 7.84, 0.80, 4.68, 5.90)
    add_pill(slide, "ФАКТИЧЕСКИЙ ВЫВОД МОДЕЛИ", 8.12, 1.09, 2.30, fill=LIME, size=7)
    add_pill(slide, "AI · AGRITECH · DRONE", 0.72, 0.65, 2.18, fill=LIME)
    add_rich_text(
        slide,
        [("Агро", WHITE, True), ("Вижн", LIME, True)],
        0.72,
        1.34,
        6.5,
        0.74,
        size=42,
    )
    add_text(slide, "Пастбище под контролем\nв реальном времени", 0.72, 2.15, 6.3, 1.45, size=31, color=WHITE, bold=True, font="Aptos Display")
    add_text(slide, "Детекция и подсчёт овец на видео с дронов, в RTSP-потоках и на загруженных кадрах.", 0.75, 3.92, 5.8, 0.74, size=15, color="C8D7E3")
    add_text(slide, "YOLO26n  ·  MPS  ·  FastAPI  ·  React  ·  PostgreSQL", 0.75, 5.19, 5.95, 0.32, size=11, color=LIME, bold=True)
    add_text(slide, "End-to-end система для агропромышленного комплекса", 0.75, 6.54, 5.8, 0.30, size=10, color="AFC2D2")
    add_footer(slide, 1, dark=True)

    # 02 — Problem
    slide = blank(prs)
    add_title(slide, "Задача", "От облёта до решения — без ручного пересчёта")
    add_picture_cover(slide, test_hero, 0.65, 1.78, 4.0, 4.98)
    add_pill(slide, "КАДР ИЗ TEST SPLIT", 0.91, 2.03, 1.72, fill=NAVY, color=WHITE, size=8)
    cards = [
        ("01", "Съёмка", "Дрон или стационарная камера дают поток с высоты."),
        ("02", "Анализ", "Модель находит каждое животное и оценивает уверенность."),
        ("03", "Контроль", "Оператор видит число голов, динамику и историю сессий."),
    ]
    for i, (num, title, body) in enumerate(cards):
        y = 1.85 + i * 1.48
        add_rect(slide, 5.03, y, 7.62, 1.16, WHITE, line="D9EBE3")
        add_text(slide, num, 5.28, y + 0.22, 0.58, 0.42, size=20, color=GREEN, bold=True)
        add_text(slide, title, 6.05, y + 0.17, 2.4, 0.3, size=16, color=INK, bold=True)
        add_text(slide, body, 6.05, y + 0.54, 5.98, 0.41, size=11.5, color=MUTED)
    add_rect(slide, 5.03, 6.30, 7.62, 0.46, NAVY)
    add_text(slide, "Один интерфейс связывает поле, модель и управленческий отчёт.", 5.28, 6.40, 7.1, 0.2, size=11, color=WHITE, bold=True)
    add_footer(slide, 2, source="Источник визуализации: локальное видео и фактический вывод модели")

    # 03 — Functionality
    slide = blank(prs, NAVY)
    add_title(slide, "Продукт", "Четыре сценария в одной системе", "Оператору не нужно переключаться между ML-скриптами, плеерами и таблицами.", dark=True)
    items = [
        ("LIVE", "Видеонаблюдение", "Два дрон-облёта одновременно; счётчик и задержка по каждому источнику.", GREEN_2),
        ("RTSP", "Подключение камеры", "URI вводится в UI, остаётся на сервере, поток восстанавливается после обрыва.", CYAN),
        ("FILE", "Фото и видео", "Загрузка материала, аннотированный результат, время обработки и уверенность.", LIME),
        ("DATA", "Отчёты", "Журнал сессий, динамика поголовья, экспорт CSV и PDF.", AMBER),
    ]
    for i, (tag, title, body, accent) in enumerate(items):
        x = 0.65 + (i % 2) * 6.1
        y = 1.95 + (i // 2) * 2.15
        add_rect(slide, x, y, 5.83, 1.72, NAVY_2, line="1C3448")
        add_pill(slide, tag, x + 0.22, y + 0.2, 0.77, fill=accent, size=8)
        add_text(slide, title, x + 1.18, y + 0.20, 4.25, 0.30, size=17, color=WHITE, bold=True)
        add_text(slide, body, x + 0.24, y + 0.76, 5.25, 0.65, size=11.5, color="B8C7D3")
    add_text(slide, "Русский UI · desktop-first · демонстрация без облака", 0.66, 6.45, 8.8, 0.27, size=11, color=LIME, bold=True)
    add_footer(slide, 3, dark=True, source="Функции реализованы в текущей сборке приложения")

    # 04 — UI overview
    slide = blank(prs)
    add_title(slide, "Интерфейс", "Главный экран показывает состояние хозяйства за секунды")
    add_frame(slide, 0.65, 1.70, 12.02, 4.92)
    add_picture_cover(slide, dashboard, 0.78, 1.83, 11.76, 4.66)
    callouts = [
        ("A", "Поголовье", 1.15, 5.81, 2.25),
        ("B", "Модель и MPS", 4.72, 5.81, 2.55),
        ("C", "Порог подсчёта", 8.30, 5.81, 2.80),
    ]
    for letter, text_value, x, y, w in callouts:
        add_rect(slide, x, y, w, 0.58, NAVY)
        add_text(slide, letter, x + 0.10, y + 0.12, 0.30, 0.25, size=10, color=LIME, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, text_value, x + 0.48, y + 0.12, w - 0.58, 0.25, size=10, color=WHITE, bold=True)
    add_footer(slide, 4, source="Скриншот: локальный интерфейс AgroVision, 12.09.2026")

    # 05 — Live & RTSP
    slide = blank(prs, NAVY)
    add_title(slide, "Видео", "Локальные облёты и RTSP в одном наблюдении", dark=True)
    add_rect(slide, 0.65, 1.78, 5.90, 3.35, NAVY_2, line="29485B")
    slide.shapes.add_movie(
        str(best_video),
        Inches(0.72),
        Inches(2.00),
        Inches(5.76),
        Inches(2.78),
        poster_frame_image=str(wide),
        mime_type="video/mp4",
    )
    add_pill(slide, "ВСТРОЕННОЕ H.264 · ПИК 1.2 С", 0.94, 4.65, 2.35, fill=LIME, size=7)
    add_picture_cover(slide, tall, 6.78, 1.78, 2.27, 3.35)
    add_rect(slide, 9.28, 1.78, 3.38, 3.35, NAVY_2, line="24445A")
    add_pill(slide, "RTSP", 9.55, 2.04, 0.78, fill=LIME, size=9)
    add_text(slide, "Камера подключается прямо из UI", 9.55, 2.58, 2.70, 0.80, size=18, color=WHITE, bold=True)
    add_bullets(
        slide,
        ["URI не возвращается клиенту", "Автопереподключение", "MJPEG-плитка на дашборде"],
        9.55,
        3.55,
        2.68,
        1.30,
        size=10.5,
        color="C0D0DC",
        accent=LIME,
    )
    add_rect(slide, 0.65, 5.42, 12.01, 1.05, "0A2631", line="1D4A55")
    steps = [("1", "Источник"), ("2", "Кадр"), ("3", "Очередь"), ("4", "YOLO26n"), ("5", "WebSocket/UI")]
    for i, (num, label) in enumerate(steps):
        x = 0.93 + i * 2.38
        add_text(slide, num, x, 5.67, 0.32, 0.30, size=12, color=LIME, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.40, 5.65, 1.55, 0.34, size=11, color=WHITE, bold=True)
        if i < 4:
            add_text(slide, "→", x + 1.92, 5.64, 0.35, 0.28, size=14, color=GREEN_2, bold=True)
    add_footer(slide, 5, dark=True, source="Встроено: аннотированный 7438342-uhd_4096_1974_30fps.mp4 · 6.7 с · H.264")

    # 06 — Threshold
    slide = blank(prs)
    add_title(slide, "Управление моделью", "Порог можно менять без перезапуска")
    add_picture_cover(slide, controls, 0.65, 1.77, 6.72, 4.78)
    add_rect(slide, 7.70, 1.77, 4.95, 4.78, WHITE, line="D9EBE3")
    add_pill(slide, "ОПЕРАЦИОННЫЙ КОНТРОЛЬ", 8.02, 2.06, 2.18, fill=NAVY, color=WHITE, size=8)
    add_text(slide, "0.40", 8.02, 2.65, 2.10, 0.80, size=38, color=GREEN, bold=True, font="Aptos Display")
    add_text(slide, "текущий порог уверенного подсчёта", 8.05, 3.40, 3.87, 0.44, size=11, color=MUTED)
    add_rect(slide, 8.04, 4.03, 3.98, 0.12, GRID, radius=False)
    add_rect(slide, 8.04, 4.03, 1.90, 0.12, GREEN, radius=False)
    add_rect(slide, 9.81, 3.92, 0.34, 0.34, LIME)
    add_bullets(
        slide,
        ["Применяется ко всем источникам", "Слабые детекции остаются видимыми", "Версия и SHA весов показаны в UI"],
        8.02,
        4.55,
        4.00,
        1.50,
        size=11.5,
    )
    add_footer(slide, 6, source="Порог ограничен сервером: 0.25–0.95")

    # 07 — Upload and examples
    slide = blank(prs, NAVY)
    add_title(slide, "Разовый анализ", "Загрузить материал — получить доказуемый результат", dark=True)
    add_picture_cover(slide, upload, 0.65, 1.77, 5.45, 2.55)
    for i, image in enumerate(examples):
        x = 6.35 + i * 2.11
        add_picture_cover(slide, image, x, 1.77, 1.87, 2.55)
        labels = [("10", "10"), ("32", "28"), ("119", "103")]
        found, truth = labels[i]
        add_rect(slide, x, 4.05, 1.87, 0.27, NAVY, radius=False, transparency=12)
        add_text(slide, f"модель {found} / разметка {truth}", x + 0.08, 4.09, 1.70, 0.13, size=6.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.65, 4.68, 12.01, 1.50, NAVY_2, line="24445A")
    flow = [
        ("01", "Валидация", "MIME · размер · пиксели"),
        ("02", "Инференс", "очередь · модель один раз"),
        ("03", "Результат", "рамки · счёт · уверенность"),
        ("04", "История", "PostgreSQL · CSV · PDF"),
    ]
    for i, (num, title, body) in enumerate(flow):
        x = 0.95 + i * 2.93
        add_text(slide, num, x, 4.99, 0.42, 0.28, size=11, color=LIME, bold=True)
        add_text(slide, title, x + 0.48, 4.95, 1.68, 0.30, size=12, color=WHITE, bold=True)
        add_text(slide, body, x + 0.48, 5.37, 2.05, 0.37, size=9.5, color="B5C7D4")
    add_footer(slide, 7, dark=True, source="Примеры: локальный test split, вывод best.pt")

    # 08 — Metrics
    slide = blank(prs)
    add_title(slide, "ML-результат", "YOLO26n уверенно находит овец на контрольной выборке", "Лучший checkpoint: эпоха 10 · изображение 640 px · inference floor 0.25")
    metrics = [("95.8", "%", "Precision", GREEN), ("93.4", "%", "Recall", CYAN), ("96.3", "%", "mAP@50", LIME), ("56.8", "%", "mAP@50–95", AMBER)]
    for i, (value, suffix, label, color) in enumerate(metrics):
        add_metric(slide, value, label, 0.65 + i * 3.02, 1.96, 2.75, color, suffix)
    add_rect(slide, 0.65, 3.38, 7.54, 2.78, WHITE, line="D9EBE3")
    add_text(slide, "Качество детекции", 0.95, 3.67, 3.0, 0.3, size=15, color=INK, bold=True)
    bars = [("Precision", 0.958, GREEN), ("Recall", 0.934, CYAN), ("mAP@50", 0.963, LIME), ("mAP@50–95", 0.568, AMBER)]
    for i, (label, value, color) in enumerate(bars):
        y = 4.18 + i * 0.43
        add_text(slide, label, 0.95, y, 1.33, 0.18, size=8.5, color=MUTED, bold=True)
        add_rect(slide, 2.34, y + 0.03, 4.76, 0.13, "E9F2EE", radius=False)
        add_rect(slide, 2.34, y + 0.03, 4.76 * value, 0.13, color, radius=False)
        add_text(slide, f"{value * 100:.1f}%", 7.16, y - 0.02, 0.65, 0.21, size=8.5, color=INK, bold=True, align=PP_ALIGN.RIGHT)
    add_rect(slide, 8.48, 3.38, 4.17, 2.78, NAVY, line=NAVY)
    add_text(slide, "Подсчёт", 8.81, 3.70, 2.8, 0.30, size=15, color=WHITE, bold=True)
    add_rich_text(slide, [("6.28", LIME, True), (" головы", WHITE, False)], 8.81, 4.28, 3.2, 0.57, size=27)
    add_text(slide, "MAE на контрольной выборке", 8.82, 4.85, 3.2, 0.28, size=10, color="B9CAD6")
    add_rich_text(slide, [("11.39", CYAN, True), (" головы", WHITE, False)], 8.81, 5.33, 3.2, 0.45, size=20)
    add_text(slide, "RMSE", 11.45, 5.42, 0.65, 0.20, size=9, color="B9CAD6", bold=True)
    add_footer(slide, 8, source="configs/model_yolo26n_aerial_sheep_v1.toml · тест содержит известную утечку соседних кадров")

    # 09 — ML lifecycle
    slide = blank(prs, NAVY)
    add_title(slide, "ML-конвейер", "Воспроизводимость от Roboflow до production-весов", dark=True)
    stages = [
        ("DATA", "Aerial Sheep v1", "Roboflow · Public Domain\nYOLO-разметка"),
        ("TRAIN", "YOLO26n", "Apple MPS · seed 42\nbest epoch 10"),
        ("EVAL", "Контроль", "P/R · mAP · MAE/RMSE\nlatency p50/p95"),
        ("SHIP", "best.pt", "20 MB · SHA-256\nfail-fast проверка"),
    ]
    for i, (tag, title, body) in enumerate(stages):
        x = 0.65 + i * 3.03
        add_rect(slide, x, 2.08, 2.62, 2.25, NAVY_2, line="26475B")
        add_pill(slide, tag, x + 0.22, 2.30, 0.82, fill=[CYAN, LIME, AMBER, GREEN_2][i], size=8)
        add_text(slide, title, x + 0.22, 2.89, 2.16, 0.36, size=16, color=WHITE, bold=True)
        add_text(slide, body, x + 0.22, 3.40, 2.16, 0.60, size=10.5, color="B5C8D5")
        if i < 3:
            add_text(slide, "→", x + 2.66, 2.94, 0.34, 0.30, size=16, color=GREEN_2, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.65, 4.74, 12.01, 1.28, "0B2633", line="1D4A55")
    add_text(slide, "SHA-256", 0.94, 5.02, 1.20, 0.25, size=10, color=LIME, bold=True)
    add_text(slide, "29561fa0c96052b…82b661d08afe5ba124e0d", 2.14, 5.00, 6.32, 0.29, size=13, color=WHITE, bold=True, font="Aptos Mono")
    add_text(slide, "Изменённые или отсутствующие веса блокируют запуск модели.", 8.66, 4.94, 3.40, 0.50, size=10.5, color="B8CAD5")
    add_footer(slide, 9, dark=True, source="Provenance и ограничения зафиксированы в model TOML и ADR")

    # 10 — Architecture
    slide = blank(prs)
    add_title(slide, "Архитектура", "Один use-case для UI, API, загрузок и потоков")
    layer_colors = [("Презентация", "React · FastAPI routes · Pydantic", "E8F8EF", GREEN), ("Приложение", "Use-cases · порты · DTO", "EAF6F7", CYAN), ("Домен", "Сущности · value objects · правила", "F5F9E6", "86B72E"), ("Инфраструктура", "YOLO · OpenCV · SQLAlchemy · RTSP", "FFF5E1", AMBER)]
    for i, (title, body, fill, accent) in enumerate(layer_colors):
        y = 1.82 + i * 1.12
        inset = i * 0.34
        add_rect(slide, 0.72 + inset, y, 6.06 - inset * 2, 0.82, fill, line=accent)
        add_text(slide, title, 0.98 + inset, y + 0.16, 1.55, 0.25, size=13, color=INK, bold=True)
        add_text(slide, body, 2.52 + inset, y + 0.18, 3.86 - inset * 2, 0.22, size=10.5, color=MUTED)
    add_rect(slide, 7.42, 1.82, 5.20, 4.19, NAVY, line=NAVY)
    add_text(slide, "Зависимости направлены внутрь", 7.76, 2.13, 4.50, 0.38, size=17, color=WHITE, bold=True)
    add_bullets(
        slide,
        [
            "Домен не знает о FastAPI и PyTorch",
            "Детектор и хранилище заменяются через порты",
            "Модель загружается один раз при старте",
            "API остаётся тонким и версионированным",
        ],
        7.76,
        2.88,
        4.36,
        2.28,
        size=12,
        color="C0D0DB",
        accent=LIME,
    )
    add_pill(slide, "МОДУЛЬНЫЙ МОНОЛИТ", 7.77, 5.35, 2.28, fill=LIME, size=8)
    add_footer(slide, 10, source="apps/api · apps/web · src/agrovision/{domain,application,infrastructure,presentation}")

    # 11 — Queue/event loop
    slide = blank(prs, NAVY)
    add_title(slide, "Runtime", "Инференс не блокирует event loop", "Все запросы проходят через ограниченную очередь к одному экземпляру модели.", dark=True)
    add_node(slide, "HTTP / RTSP / файл", "асинхронные источники", 0.65, 2.24, 2.32, 1.18, fill=NAVY_2, accent=CYAN, dark=True)
    add_node(slide, "Bounded queue", "backpressure и порядок", 3.57, 2.24, 2.32, 1.18, fill=NAVY_2, accent=LIME, dark=True)
    add_node(slide, "Worker thread", "PyTorch вне asyncio", 6.49, 2.24, 2.32, 1.18, fill=NAVY_2, accent=AMBER, dark=True)
    add_node(slide, "YOLO26n", "одна модель · MPS/CPU", 9.41, 2.24, 2.32, 1.18, fill=NAVY_2, accent=GREEN_2, dark=True)
    for x in (2.98, 5.90, 8.82):
        connect(slide, x, 2.83, x + 0.57, 2.83, color=GREEN_2, width=2.5)
    add_rect(slide, 1.26, 4.34, 10.79, 1.28, "0B2633", line="1F4D59")
    facts = [("21 мс", "задержка на дашборде"), ("MPS", "локально на Apple Silicon"), ("CPU", "в Linux-контейнере"), ("1×", "checkpoint в памяти")]
    for i, (value, label) in enumerate(facts):
        x = 1.56 + i * 2.62
        add_text(slide, value, x, 4.62, 1.20, 0.40, size=20, color=[CYAN, LIME, AMBER, GREEN_2][i], bold=True)
        add_text(slide, label, x, 5.05, 1.95, 0.28, size=9.5, color="B8CAD5")
    add_footer(slide, 11, dark=True, source="Runtime реализован через QueuedDetector и run_in_executor")

    # 12 — Stack
    slide = blank(prs)
    add_title(slide, "Технологии", "Стек собран вокруг быстрой разработки и понятного production-пути")
    columns = [
        ("ML / VIDEO", [("Python 3.12", "обучение и сервис"), ("YOLO26n", "детекция"), ("PyTorch + MPS", "локальное ускорение"), ("OpenCV", "видео и RTSP")], GREEN),
        ("BACKEND / DATA", [("FastAPI", "REST + WebSocket"), ("Pydantic v2", "контракты"), ("SQLAlchemy async", "репозитории"), ("PostgreSQL 16", "история и пользователи")], CYAN),
        ("FRONTEND / OPS", [("React + TypeScript", "интерфейс"), ("Tailwind + Recharts", "визуализация"), ("Docker Compose", "полный стек"), ("Nginx", "SPA и reverse proxy")], AMBER),
    ]
    for i, (heading, values, accent) in enumerate(columns):
        x = 0.65 + i * 4.05
        add_rect(slide, x, 1.84, 3.76, 4.56, WHITE, line="D9EBE3")
        add_pill(slide, heading, x + 0.23, 2.08, 1.62, fill=accent, size=8)
        for j, (name, use) in enumerate(values):
            y = 2.78 + j * 0.79
            add_text(slide, name, x + 0.24, y, 2.08, 0.27, size=13, color=INK, bold=True)
            add_text(slide, use, x + 0.24, y + 0.34, 2.92, 0.20, size=9.5, color=MUTED)
            if j < 3:
                add_rect(slide, x + 0.24, y + 0.66, 3.26, 0.01, GRID, radius=False)
    add_footer(slide, 12, source="Dependency lock: uv.lock · frontend lock: apps/web/package-lock.json")

    # 13 — Docker
    slide = blank(prs, NAVY)
    add_title(slide, "Развёртывание", "Docker Compose поднимает web, API и PostgreSQL", dark=True)
    add_node(slide, "Браузер", "http://localhost:8080", 0.65, 2.44, 1.90, 1.16, fill=NAVY_2, accent=LIME, dark=True)
    add_node(slide, "Web + Nginx", "SPA · reverse proxy", 3.14, 2.44, 2.14, 1.16, fill=NAVY_2, accent=CYAN, dark=True)
    add_node(slide, "FastAPI", "CPU · port 8000", 5.87, 2.44, 2.14, 1.16, fill=NAVY_2, accent=GREEN_2, dark=True)
    add_node(slide, "PostgreSQL 16", "volume · healthcheck", 8.60, 2.44, 2.45, 1.16, fill=NAVY_2, accent=AMBER, dark=True)
    for x in (2.56, 5.29, 8.02):
        connect(slide, x, 3.02, x + 0.56, 3.02, color=GREEN_2, width=2.5)
    add_rect(slide, 11.47, 2.44, 1.20, 1.16, "102C38", line="315669")
    add_text(slide, "VOL", 11.64, 2.70, 0.84, 0.25, size=11, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "данные", 11.58, 3.12, 0.96, 0.18, size=8, color="BCCCD6", align=PP_ALIGN.CENTER)
    add_rect(slide, 0.65, 4.38, 12.02, 1.42, "0B2633", line="1F4D59")
    docker_facts = [
        ("READ-ONLY", "best.pt, demo-видео, результаты"),
        ("FAIL-FAST", "без секретов compose не стартует"),
        ("NON-ROOT", "API работает под uid 10001"),
        ("HEALTH", "DB → API → Web"),
    ]
    for i, (title, body) in enumerate(docker_facts):
        x = 0.95 + i * 2.93
        add_text(slide, title, x, 4.70, 1.35, 0.26, size=10, color=[LIME, CYAN, GREEN_2, AMBER][i], bold=True)
        add_text(slide, body, x, 5.12, 2.28, 0.36, size=9.5, color="BBCBD6")
    add_text(slide, "make configure  →  make docker-up", 0.68, 6.24, 6.5, 0.34, size=15, color=WHITE, bold=True, font="Aptos Mono")
    add_footer(slide, 13, dark=True, source="Локальный запуск: MPS + SQLite · Docker: CPU + PostgreSQL")

    # 14 — Security
    slide = blank(prs)
    add_title(slide, "Безопасность", "Секреты не хранятся в коде, данные проходят проверку")
    add_rect(slide, 0.65, 1.75, 4.02, 4.87, NAVY, line=NAVY)
    add_pill(slide, "ГДЕ ЛЕЖАТ СЕКРЕТЫ", 0.96, 2.04, 1.88, fill=LIME, size=8)
    add_text(slide, ".env", 0.96, 2.72, 2.4, 0.62, size=33, color=WHITE, bold=True, font="Aptos Mono")
    add_text(slide, "локальная разработка", 0.98, 3.34, 2.6, 0.25, size=10, color="B8CAD5")
    add_bullets(
        slide,
        ["права файла 0600", "исключён через .gitignore", "значения не печатаются в лог", "в контейнер передаются через env"],
        0.96,
        3.94,
        3.15,
        1.85,
        size=11,
        color="C4D2DC",
        accent=LIME,
    )
    security = [
        ("AUTH", "Argon2 для паролей · JWT access/refresh"),
        ("UPLOAD", "MIME · размер до 100 MB · пиксели · безопасные имена"),
        ("MODEL", "SHA-256 checkpoint · fallback отключён"),
        ("NETWORK", "CORS · security headers · RTSP URI server-side"),
        ("RUNTIME", "bounded queue · non-root container · read-only mounts"),
    ]
    for i, (tag, body) in enumerate(security):
        y = 1.75 + i * 0.92
        add_pill(slide, tag, 5.03, y + 0.10, 0.88, fill=[GREEN, CYAN, AMBER, GREEN_2, LIME][i], color=NAVY, size=7)
        add_text(slide, body, 6.14, y + 0.07, 6.02, 0.45, size=12.5, color=INK, bold=True)
        if i < 4:
            add_rect(slide, 5.03, y + 0.77, 7.28, 0.01, GRID, radius=False)
    add_rect(slide, 5.03, 6.15, 7.61, 0.47, "FFF3D4", line="F4D48B")
    add_text(slide, "Production: TLS · secret manager · HttpOnly SameSite cookies · закрытый порт БД", 5.25, 6.27, 7.15, 0.18, size=9.5, color="795319", bold=True)
    add_footer(slide, 14, source="Проверено: .env mode 0600; значения секретов в презентацию не включены")

    # 15 — Honest limits and close
    slide = blank(prs, NAVY)
    add_picture_cover(slide, wide, 7.54, 0, 5.79, 7.5, brightness=0.42)
    add_rect(slide, 6.80, 0, 6.53, 7.5, NAVY, radius=False, transparency=30)
    add_picture_cover(slide, wide, 7.62, 0.58, 5.10, 6.34)
    add_pill(slide, "39 ОВЕЦ В КАДРЕ", 8.01, 0.94, 1.47, fill=LIME, size=8)
    add_pill(slide, "ГОТОВО К ДЕМО", 0.72, 0.66, 1.62, fill=LIME, size=8)
    add_text(slide, "АгроВижн превращает\nвидеопоток в действие", 0.72, 1.40, 6.25, 1.25, size=30, color=WHITE, bold=True, font="Aptos Display")
    add_text(slide, "Фото · видео · два дрон-облёта · RTSP · дашборд · отчёты", 0.74, 2.90, 5.95, 0.55, size=13, color="C2D1DC")
    add_rect(slide, 0.72, 3.80, 5.98, 1.48, "0B2633", line="1D4A55")
    add_text(slide, "Честное ограничение", 1.02, 4.08, 2.45, 0.28, size=12, color=AMBER, bold=True)
    add_text(slide, "Соседние кадры исходных DJI-видео попали в разные splits, поэтому метрики оптимистичны. Следующий шаг — split по полётам и полевое тестирование.", 1.02, 4.52, 5.20, 0.52, size=10.5, color="C3D2DB")
    add_text(slide, "Давайте покажем живой сценарий", 0.74, 5.82, 5.35, 0.46, size=20, color=LIME, bold=True)
    add_text(slide, "localhost:5173  ·  operator demo", 0.75, 6.47, 4.8, 0.26, size=10, color="AFC2D1", font="Aptos Mono")
    add_footer(slide, 15, dark=True)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(DECK_PATH)
    return DECK_PATH


if __name__ == "__main__":
    path = build()
    print(f"created {path.relative_to(ROOT)}")
